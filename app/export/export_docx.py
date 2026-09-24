"""
export_docx.py
================
Génère un rapport Word (.docx) à partir du rapport JSON final
(data/processed/rapport_final.json), avec les sections narratives et un
tableau récapitulatif des KPIs par section.

Usage :
    python app/export/export_docx.py --in data/processed/rapport_final.json --out rapport_esg.docx
"""

import argparse
import json
import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def construire_document(rapport: dict) -> Document:
    doc = Document()

    titre = doc.add_heading("Rapport de durabilité — NordTech Industries", level=0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sous_titre = doc.add_paragraph(f"Exercice {rapport['annee']} — aligné sur la taxonomie ESRS")
    sous_titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous_titre.runs[0].italic = True

    note = doc.add_paragraph(
        "Document généré automatiquement à partir de données ESG fictives, à des fins "
        "de démonstration (projet portfolio). Chaque chiffre cité a fait l'objet d'un "
        "contrôle automatisé de correspondance avec les données source."
    )
    note.runs[0].font.size = Pt(9)
    note.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    for section in rapport["sections"]:
        doc.add_heading(section["esrs_section"], level=1)

        statut = section["statut_final"]
        p_statut = doc.add_paragraph()
        run = p_statut.add_run("✅ Vérifié" if statut == "verifie" else "⚠️ À corriger — revue humaine requise")
        run.bold = True
        run.font.color.rgb = RGBColor(0x15, 0x87, 0x24) if statut == "verifie" else RGBColor(0x85, 0x64, 0x04)

        doc.add_paragraph(section["texte_final"])

        # Tableau récapitulatif des KPIs de la section
        kpis = section.get("kpis_utilises", [])
        if kpis:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            for i, h in enumerate(["Indicateur", "Valeur", "Unité", "Variation YoY", "Statut vs objectif"]):
                hdr[i].text = h
            for kpi in kpis:
                row = table.add_row().cells
                row[0].text = str(kpi.get("nom_indicateur", ""))
                row[1].text = str(kpi.get("valeur", ""))
                row[2].text = str(kpi.get("unite", ""))
                variation = kpi.get("variation_pct")
                row[3].text = f"{variation:+.2f}%" if variation not in (None, "nan", "") else "n/a"
                row[4].text = str(kpi.get("statut_vs_objectif", ""))

        doc.add_paragraph()  # espacement

    return doc


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="input_path",
                         default=os.path.join("data", "processed", "rapport_final.json"))
    parser.add_argument("--out", dest="output_path",
                         default=os.path.join("data", "processed", "rapport_esg.docx"))
    args = parser.parse_args()

    with open(args.input_path, encoding="utf-8") as f:
        rapport = json.load(f)

    doc = construire_document(rapport)
    doc.save(args.output_path)
    print(f"[OK] Rapport Word généré : {args.output_path}")


if __name__ == "__main__":
    main()
